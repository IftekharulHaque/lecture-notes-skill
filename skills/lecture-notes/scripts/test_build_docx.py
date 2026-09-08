#!/usr/bin/env python3
"""Self-check for build_docx.py. Run: python test_build_docx.py

Builds a doc from a spec and asserts the styled structure survives the roundtrip.
No test framework — just asserts, per the repo's keep-it-lazy rule.
"""
import os
import tempfile

from docx import Document

from build_docx import build

SPEC = {
    "title": "Test Lecture",
    "date": "October 2025",
    "overview": "Covers testing.",
    "key_topics": [{"term": "Stub", "gloss": "canned data"}],
    "notes": [
        {"type": "heading", "level": 2, "text": "Section A"},
        {"type": "paragraph", "text": "Body text."},
        {"type": "bullets", "items": ["first", "second"]},
    ],
    "todos": ["Read chapter 6"],
}


def _styles(doc):
    return [(p.style.name, p.text) for p in doc.paragraphs]


def test_full_spec():
    with tempfile.TemporaryDirectory() as d:
        out = build(SPEC, os.path.join(d, "out.docx"))
        got = _styles(Document(out))

    styles = {s for s, _ in got}
    texts = [t for _, t in got]
    assert ("Heading 1", "Test Lecture") in got, "title not Heading 1"
    assert "Heading 2" in styles, "section/overview headings missing"
    assert any(s == "List Bullet" and t == "first" for s, t in got), "bullets not styled"
    assert any(t == "Stub — canned data" for t in texts), "key topic gloss missing"
    assert any(t.startswith("☐ ") and "chapter 6" in t for t in texts), "todo checkbox missing"


def test_empty_todos_renders_fallback_line():
    spec = {"title": "Bare", "todos": []}
    with tempfile.TemporaryDirectory() as d:
        out = build(spec, os.path.join(d, "out.docx"))
        texts = [p.text for p in Document(out).paragraphs]
    assert "No action items mentioned in this lecture." in texts, "empty-todos fallback missing"


if __name__ == "__main__":
    test_full_spec()
    test_empty_todos_renders_fallback_line()
    print("ok")
